from pathlib import Path

from pptx import Presentation


class PresentationGenerator:
    def generate(self, topic, slides, user_profile):
        presentation = Presentation()
        title_slide_layout = presentation.slide_layouts[0]
        slide = presentation.slides.add_slide(title_slide_layout)
        slide.shapes.title.text = topic
        slide.placeholders[1].text = f"Профиль: {user_profile.get('specialty', 'не указан')}"
        for index in range(1, slides):
            layout = presentation.slide_layouts[1]
            slide = presentation.slides.add_slide(layout)
            slide.shapes.title.text = f"Слайд {index + 1}"
            slide.placeholders[1].text = "TODO: контент с GPT-4"
        output_dir = Path("media/presentations")
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / f"{topic.replace(' ', '_')}.pptx"
        presentation.save(output_path)
        return {"file_url": f"/media/presentations/{output_path.name}", "slides": slides}

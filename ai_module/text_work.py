from pathlib import Path

from docx import Document


class TextWorkGenerator:
    def generate(self, topic, pages, user_profile):
        document = Document()
        document.add_heading(topic, level=1)
        document.add_paragraph(
            f"Профиль студента: {user_profile.get('specialty', 'не указан')}"
        )
        for index in range(pages):
            document.add_paragraph(
                f"Раздел {index + 1}. TODO: сгенерировать текст через GPT-4."
            )
        output_dir = Path("media/works")
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / f"{topic.replace(' ', '_')}.docx"
        document.save(output_path)
        return {"file_url": f"/media/works/{output_path.name}", "pages": pages}
